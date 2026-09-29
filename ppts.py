import copy
from pptx import Presentation



LAYOUT_FILES = {
    1: "bodytemplates/3 points/1.pptx",
    2: "bodytemplates/3 points/2.pptx",
    3: "bodytemplates/4 points/3.pptx",
    4: "bodytemplates/4 points/4.pptx",
    5: "bodytemplates/4 points/5.pptx",
}

INTRO_PATH = "introtemplates/introt1.pptx"
END_PATH = "endtemplates/endt1.pptx"



def get_intro_slide(intro_title: str):
    """Loads the intro template and fills in the title."""
    template_prs = Presentation(INTRO_PATH)
    slide = template_prs.slides[0]

    for shape in slide.shapes:
        if shape.name == "title":
            shape.text_frame.text = intro_title

    return slide


def get_body_slide(layout_number: int, subtopic: str, bullets: list):
    """Loads the matching body template and fills in subtopic + bullets."""
    path = LAYOUT_FILES[layout_number]
    template_prs = Presentation(path)
    slide = template_prs.slides[0]

    for shape in slide.shapes:
        if shape.name == "subtopic":
            shape.text_frame.text = subtopic
        for idx, b in enumerate(bullets, start=1):
            if shape.name == f"head{idx}":
                shape.text_frame.text = b.bullethead
            elif shape.name == f"content{idx}":
                shape.text_frame.text = b.bulletcontent

    return slide


def get_end_slide(ending_line: str):
    """Loads the end template and fills in the closing line."""
    template_prs = Presentation(END_PATH)
    slide = template_prs.slides[0]

    for shape in slide.shapes:
        if shape.name == "ending":
            shape.text_frame.text = ending_line

    return slide



def append_slide(final_prs, slide):
    """Copies all shapes from `slide` into a new slide in final_prs."""
    new_slide = final_prs.slides.add_slide(final_prs.slide_layouts[6])
    for shape in slide.shapes:
        new_slide.shapes._spTree.append(copy.deepcopy(shape.element))
    return new_slide



def generate_ppt(topic, subtopics, contents, layouts):
    final_prs = Presentation()
    final_prs.slide_width = 12192000
    final_prs.slide_height = 6858000

    intro_slide = get_intro_slide(topic)
    append_slide(final_prs, intro_slide)

    for i, subtopic in enumerate(subtopics):
        bullets = contents[i].bullets
        layout_num = layouts[i]

        body_slide = get_body_slide(layout_num, subtopic, bullets)
        append_slide(final_prs, body_slide)


    ending_line = f"Thank you for exploring {topic}."
    end_slide = get_end_slide(ending_line)
    append_slide(final_prs, end_slide)

    final_prs.save("final_output.pptx")
    return final_prs